"""
data_loader.py
Carga y procesa los archivos Excel de coberturas de vacunación DIRESA Huancavelica 2026
"""

import pandas as pd
import openpyxl
import io
from datetime import datetime


# ── Semáforo de cobertura ──────────────────────────────────────────────
SEMAPHORE_THRESHOLDS = [
    (0,      20.0,   "Crítico",    "#F44336", "#FF7043"),
    (20.0,   23.75,  "En Proceso", "#FFC107", "#FFD54F"),
    (23.75,  200.0,  "Óptimo",     "#4CAF50", "#81C784"),
]

def apply_semaphore(pct: float) -> dict:
    """Retorna dict con label, color_fill y color_light para un % dado."""
    for lo, hi, label, color, light in SEMAPHORE_THRESHOLDS:
        if lo <= pct < hi:
            return {"label": label, "color": color, "light": light}
    # >= 200 (solo por seguridad)
    return {"label": "Óptimo", "color": "#4CAF50", "light": "#81C784"}


# ── Mapeo de vacunas por hoja ──────────────────────────────────────────
VACUNAS_INFO = {
    "3° Dosis IPV": {
        "sheet_idx": 0,
        "col_meta": 1,       # columna B → META (< 1 año)
        "col_vac": 2,        # columna C → Nro. vacunados
        "edad": "< 1 Año",
        "skip_rows": 2,      # saltar 2 filas de encabezado
    },
    "3° Dosis Pentavalente": {
        "sheet_idx": 0,
        "col_meta": 1,
        "col_vac": 4,        # columna E
        "edad": "< 1 Año",
        "skip_rows": 2,
    },
    "1° Ref DPT": {
        "sheet_idx": 1,
        "col_meta": 1,       # columna B → META 1 AÑO
        "col_vac": 2,        # columna C
        "edad": "1 Año",
        "skip_rows": 4,
    },
    "2° Ref DPT": {
        "sheet_idx": 1,
        "col_meta": 4,       # columna E → META 4 AÑOS
        "col_vac": 5,        # columna F
        "edad": "4 Años",
        "skip_rows": 4,
    },
    "1° Dosis SPR": {
        "sheet_idx": 2,
        "col_meta": 1,
        "col_vac": 2,
        "edad": "1 Año",
        "skip_rows": 3,
    },
    "2° Dosis SPR": {
        "sheet_idx": 2,
        "col_meta": 1,
        "col_vac": 4,
        "edad": "1 Año",
        "skip_rows": 3,
    },
}

PROVINCIAS = [
    "ACOBAMBA", "ANGARAES", "CASTROVIRREYNA",
    "CHURCAMPA", "HUANCAVELICA", "HUAYTARA", "TAYACAJA",
]


def _read_sheet_data(wb, sheet_idx: int, col_meta: int, col_vac: int, skip_rows: int) -> dict:
    """Lee una hoja y retorna dict {provincia: (vacunados, meta)}."""
    ws = wb.worksheets[sheet_idx]
    result = {}
    for row in ws.iter_rows(min_row=skip_rows + 1, values_only=True):
        prov = str(row[0]).strip().upper() if row[0] else ""
        if prov in PROVINCIAS:
            try:
                meta = float(row[col_meta]) if row[col_meta] else 0
                vac  = float(row[col_vac])  if row[col_vac]  else 0
                result[prov] = (vac, meta)
            except (TypeError, ValueError):
                result[prov] = (0, 0)
    return result


def load_coverage_data(file_source) -> pd.DataFrame:
    """
    Carga el Excel extendido de coberturas.
    file_source: path (str) o BytesIO (upload Streamlit).
    Retorna DataFrame con columnas:
        provincia, vacuna, edad, vacunados, meta, cobertura_pct,
        semaforo_label, semaforo_color, semaforo_light
    """
    if isinstance(file_source, (str, bytes)):
        wb = openpyxl.load_workbook(file_source, data_only=True)
    else:
        data = file_source.read()
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)

    rows = []
    for vacuna, cfg in VACUNAS_INFO.items():
        sheet_data = _read_sheet_data(
            wb,
            cfg["sheet_idx"],
            cfg["col_meta"],
            cfg["col_vac"],
            cfg["skip_rows"],
        )
        for prov in PROVINCIAS:
            vac, meta = sheet_data.get(prov, (0, 0))
            pct = (vac / meta * 100) if meta > 0 else 0.0
            sem = apply_semaphore(pct)
            rows.append({
                "provincia":       prov,
                "vacuna":          vacuna,
                "edad":            cfg["edad"],
                "vacunados":       int(vac),
                "meta":            int(meta),
                "cobertura_pct":   round(pct, 2),
                "semaforo_label":  sem["label"],
                "semaforo_color":  sem["color"],
                "semaforo_light":  sem["light"],
            })

    return pd.DataFrame(rows)


def get_regional_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Resumen regional (todas provincias sumadas) por vacuna."""
    agg = df.groupby("vacuna").agg(
        vacunados=("vacunados", "sum"),
        meta=("meta", "sum"),
        edad=("edad", "first"),
    ).reset_index()
    agg["cobertura_pct"] = (agg["vacunados"] / agg["meta"] * 100).round(2)
    agg["cobertura_pct"] = agg["cobertura_pct"].fillna(0)
    sem = agg["cobertura_pct"].apply(apply_semaphore)
    agg["semaforo_label"] = sem.apply(lambda x: x["label"])
    agg["semaforo_color"] = sem.apply(lambda x: x["color"])
    return agg


def get_ris_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Tabla pivote: filas=provincia, columnas=vacuna, valores=cobertura%."""
    pivot = df.pivot_table(
        index="provincia",
        columns="vacuna",
        values="cobertura_pct",
        aggfunc="first",
    ).reset_index()
    pivot = pivot[["provincia"] + list(VACUNAS_INFO.keys())]
    return pivot


def get_phva_analysis(df: pd.DataFrame) -> dict:
    """Genera análisis PHVA-Deming."""
    total_meta     = df["meta"].sum()
    total_vacunados = df["vacunados"].sum()
    cobertura_gral = round(total_vacunados / total_meta * 100, 2) if total_meta > 0 else 0

    criticos   = df[df["semaforo_label"] == "Crítico"]["provincia"].unique().tolist()
    en_proceso = df[df["semaforo_label"] == "En Proceso"]["provincia"].unique().tolist()
    optimos    = df[df["semaforo_label"] == "Óptimo"]["provincia"].unique().tolist()

    por_vacuna = df.groupby("vacuna").agg(
        vacunados=("vacunados", "sum"),
        meta=("meta", "sum"),
    ).assign(
        cobertura_pct=lambda x: (x["vacunados"] / x["meta"] * 100).round(2)
    ).reset_index()

    vacuna_mas_baja  = por_vacuna.loc[por_vacuna["cobertura_pct"].idxmin(), "vacuna"]
    vacuna_mas_alta  = por_vacuna.loc[por_vacuna["cobertura_pct"].idxmax(), "vacuna"]
    pct_mas_baja     = por_vacuna["cobertura_pct"].min()
    pct_mas_alta     = por_vacuna["cobertura_pct"].max()

    brecha_total = total_meta - total_vacunados

    return {
        "cobertura_gral":    cobertura_gral,
        "total_meta":        total_meta,
        "total_vacunados":   total_vacunados,
        "brecha_total":      brecha_total,
        "provincias_criticas":   criticos,
        "provincias_en_proceso": en_proceso,
        "provincias_optimas":    optimos,
        "vacuna_mas_baja":   vacuna_mas_baja,
        "pct_mas_baja":      pct_mas_baja,
        "vacuna_mas_alta":   vacuna_mas_alta,
        "pct_mas_alta":      pct_mas_alta,
        "por_vacuna":        por_vacuna,
        "fecha_analisis":    datetime.now().strftime("%d/%m/%Y %H:%M"),
    }
